
import time
from pathlib import Path

import boto3
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# CONFIGURATION
# --------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REPORT_DIR = PROJECT_ROOT / "reports"
REPORT_FILE = REPORT_DIR / "compliance_gold_report.xlsx"

AWS_REGION = "us-east-1"
DATABASE = "compliance_analytics"
WORKGROUP = "primary"

ATHENA_OUTPUT = (
    "s3://amandataeng-compliance-de-2026/athena-results/"
)

athena = boto3.client("athena", region_name=AWS_REGION)


# --------------------------------------------------
# ATHENA QUERY HELPER
# --------------------------------------------------

def run_query(sql):
    response = athena.start_query_execution(
        QueryString=sql,
        QueryExecutionContext={"Database": DATABASE},
        WorkGroup=WORKGROUP,
    )

    query_id = response["QueryExecutionId"]

    while True:
        result = athena.get_query_execution(
            QueryExecutionId=query_id
        )

        status = result["QueryExecution"]["Status"]["State"]

        if status == "SUCCEEDED":
            break

        if status in ("FAILED", "CANCELLED"):
            reason = result["QueryExecution"]["Status"].get(
                "StateChangeReason", "Unknown error"
            )
            raise RuntimeError(
                f"Athena query {status}: {reason}"
            )

        time.sleep(2)

    rows = []
    paginator = athena.get_paginator("get_query_results")

    for page in paginator.paginate(
        QueryExecutionId=query_id
    ):
        rows.extend(
            page["ResultSet"]["Rows"]
        )

    if not rows:
        return pd.DataFrame()

    headers = [
        cell.get("VarCharValue", "")
        for cell in rows[0]["Data"]
    ]

    records = []

    # First row contains column names, so skip it.
    for row in rows[1:]:
        values = [
            cell.get("VarCharValue")
            for cell in row.get("Data", [])
        ]

        values += [None] * (len(headers) - len(values))
        records.append(values[:len(headers)])

    return pd.DataFrame(records, columns=headers)


# --------------------------------------------------
# READ EXISTING GOLD TABLES
# --------------------------------------------------

def load_gold_data():
    print("Reading existing Gold tables from Athena...")

    kpi_df = run_query("""
        SELECT kpi_status, total_device, percentage
        FROM gold_kpi_summary_live
        ORDER BY kpi_status
    """)

    control_df = run_query("""
        SELECT control_name, control_status,
               total_devices, percentage
        FROM gold_summary_live
        ORDER BY control_name, control_status
    """)

    if kpi_df.empty or control_df.empty:
        raise ValueError(
            "A Gold table returned no data. "
            "Check the existing Athena tables."
        )

    kpi_df["total_device"] = pd.to_numeric(
        kpi_df["total_device"], errors="raise"
    )

    kpi_df["percentage"] = pd.to_numeric(
        kpi_df["percentage"], errors="raise"
    )

    control_df["total_devices"] = pd.to_numeric(
        control_df["total_devices"], errors="raise"
    )

    control_df["percentage"] = pd.to_numeric(
        control_df["percentage"], errors="raise"
    )

    return kpi_df, control_df


# --------------------------------------------------
# EXCEL FORMATTING HELPERS
# --------------------------------------------------

def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = PatternFill(
                "solid", fgColor="1F4E78"
            )
            cell.font = Font(
                color="FFFFFF", bold=True
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


def format_data_sheet(ws, percentage_columns=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    style_header(ws)

    for column_cells in ws.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 3, 14), 32
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in percentage_columns:
                # Gold percentages are already values like 80.33,
                # not decimal fractions like 0.8033.
                cell.number_format = '0.00"%"'


def add_card(ws, cell_range, title, value):
    ws.merge_cells(cell_range)

    cell = ws[cell_range.split(":")[0]]
    cell.value = f"{title}\n{value}"
    cell.fill = PatternFill(
        "solid", fgColor="D9EAF7"
    )
    cell.font = Font(
        name="Calibri", size=14, bold=True,
        color="17365D"
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    for row in ws[cell_range]:
        for item in row:
            item.fill = PatternFill(
                "solid", fgColor="D9EAF7"
            )
            item.border = Border(
                bottom=Side(style="thin", color="9EADBA")
            )


# --------------------------------------------------
# BUILD DASHBOARD
# --------------------------------------------------


def create_dashboard(wb, kpi_df, control_df):
    from openpyxl.chart import PieChart, BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False

    # Consistent dashboard spacing
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # Title
    ws.merge_cells("A1:P1")
    ws["A1"] = "IT ASSET COMPLIANCE DASHBOARD"
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    # KPI cards
    lookup = {
        str(r["kpi_status"]).strip().lower(): int(r["total_device"])
        for _, r in kpi_df.iterrows()
    }

    cards = [
        ("A3:D5", "TOTAL DEVICES", int(kpi_df["total_device"].sum())),
        ("E3:H5", "COMPLIANT", lookup.get("compliant", 0)),
        ("I3:L5", "NON-COMPLIANT", lookup.get("non-compliant", 0)),
        ("M3:P5", "EXCLUDED", lookup.get("excluded", 0)),
    ]

    for area, label, value in cards:
        ws.merge_cells(area)
        cell = ws[area.split(":")[0]]
        cell.value = f"{label}\n{value:,}"
        cell.font = Font(size=15, bold=True, color="17365D")
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        for row in ws[area]:
            for c in row:
                c.fill = PatternFill("solid", fgColor="D9EAF7")
                c.border = Border(
                    bottom=Side(style="thin", color="9EADBA")
                )

    for r in range(3, 6):
        ws.row_dimensions[r].height = 25

    # Separate hidden chart-data sheet
    if "Chart Data" in wb.sheetnames:
        del wb["Chart Data"]

    data = wb.create_sheet("Chart Data")
    data.append(["KPI Status", "Device Count"])

    for _, r in kpi_df.iterrows():
        data.append([str(r["kpi_status"]), int(r["total_device"])])

    data.append([])
    control_header = len(kpi_df) + 3
    data.cell(control_header, 1, "Control / Status")
    data.cell(control_header, 2, "Device Count")

    for i, (_, r) in enumerate(control_df.iterrows(), start=control_header + 1):
        data.cell(i, 1, f"{r['control_name']} - {r['control_status']}")
        data.cell(i, 2, int(r["total_devices"]))

    data.sheet_state = "hidden"

    # Pie chart: only category and percentage labels
    pie = PieChart()
    pie.title = "Overall KPI Distribution"
    pie.style = 10
    pie.height = 9
    pie.width = 14
    pie.varyColors = True

    pie.add_data(
        Reference(data, min_col=2, min_row=1, max_row=len(kpi_df) + 1),
        titles_from_data=True,
    )
    pie.set_categories(
        Reference(data, min_col=1, min_row=2, max_row=len(kpi_df) + 1)
    )

    pie.dataLabels = DataLabelList()
    pie.dataLabels.showLegendKey = False
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = True
    pie.dataLabels.showSerName = False
    pie.dataLabels.showPercent = True
    pie.dataLabels.showLeaderLines = True
    pie.legend.position = "r"

    ws.add_chart(pie, "A7")

    # Horizontal bar chart with category names visible on the left
    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Device Count by Control and Status"
    bar.height = 15
    bar.width = 23
    bar.legend = None
    bar.gapWidth = 55
    bar.overlap = 0

    control_end = control_header + len(control_df)

    bar.add_data(
        Reference(
            data, min_col=2,
            min_row=control_header,
            max_row=control_end,
        ),
        titles_from_data=True,
    )
    bar.set_categories(
        Reference(
            data, min_col=1,
            min_row=control_header + 1,
            max_row=control_end,
        )
    )

    # For horizontal bars, x_axis is the category axis.
    bar.x_axis.title = "Control / Status"
    bar.y_axis.title = "Number of Devices"
    bar.x_axis.tickLblPos = "low"
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.y_axis.scaling.orientation = "maxMin"

    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.showCatName = False
    bar.dataLabels.showSerName = False
    bar.dataLabels.position = "outEnd"

    ws.add_chart(bar, "H7")

    # Place summary well below the chart
    title_row = 38
    ws.merge_cells(start_row=title_row, start_column=1,
                   end_row=title_row, end_column=8)

    title = ws.cell(title_row, 1, "CONTROL STATUSES NEEDING ATTENTION")
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    ws.row_dimensions[title_row].height = 24

    for col, heading in enumerate(["Control", "Status", "Devices"], start=1):
        c = ws.cell(title_row + 1, col, heading)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")

    attention = {
        "bitlocker": {"disabled", "unknown"},
        "antivirus": {"outdated", "missing"},
        "patch": {"outdated", "pending", "unknown"},
        "dlp": {"non-compliant", "unknown"},
        "proxy": {"missing", "unknown"},
    }

    out_row = title_row + 2
    for _, r in control_df.iterrows():
        control = str(r["control_name"]).strip().lower()
        status = str(r["control_status"]).strip().lower()

        if status in attention.get(control, set()):
            ws.cell(out_row, 1, r["control_name"])
            ws.cell(out_row, 2, r["control_status"])
            ws.cell(out_row, 3, int(r["total_devices"]))
            out_row += 1

    note_row = out_row + 2
    ws.merge_cells(
        start_row=note_row, start_column=1,
        end_row=note_row + 1, end_column=14
    )
    note = ws.cell(
        note_row, 1,
        "Note: Counts are grouped by control and status. A device may "
        "appear under multiple controls, so these counts do not represent "
        "unique devices. Flagged statuses require review and are not all "
        "confirmed non-compliance cases."
    )
    note.font = Font(size=10, italic=True, color="666666")
    note.alignment = Alignment(wrap_text=True, vertical="top")

    # Better screen and print layout
    ws.freeze_panes = "A7"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_area = f"A1:P{note_row + 1}"

# --------------------------------------------------
# MAIN
# --------------------------------------------------

def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    kpi_df, control_df = load_gold_data()

    wb = Workbook()

    # Write the existing summary sheets.
    kpi_ws = wb.active
    kpi_ws.title = "KPI Summary"

    for row in [list(kpi_df.columns)] + kpi_df.values.tolist():
        kpi_ws.append(row)

    control_ws = wb.create_sheet("Control Summary")

    for row in [
        list(control_df.columns)
    ] + control_df.values.tolist():
        control_ws.append(row)

    format_data_sheet(kpi_ws, percentage_columns=(3,))
    format_data_sheet(control_ws, percentage_columns=(4,))

    # Add the dashboard as the first sheet.
    create_dashboard(wb, kpi_df, control_df)

    wb.active = 0
    wb.save(REPORT_FILE)

    print("\nEXCEL DASHBOARD CREATED SUCCESSFULLY")
    print(f"Total devices: {int(kpi_df['total_device'].sum()):,}")
    print(f"KPI rows: {len(kpi_df)}")
    print(f"Control summary rows: {len(control_df)}")
    print(f"Output: {REPORT_FILE}")


if __name__ == "__main__":
    main()
